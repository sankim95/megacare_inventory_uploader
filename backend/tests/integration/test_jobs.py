from __future__ import annotations

import hashlib
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import func, select

from app.core.config import Settings, get_settings
from app.core.database import Base, build_engine, build_session_factory, get_db
from app.main import create_app
from app.models import (
    CompletedDocument,
    Document,
    Item,
    Job,
    JobStatus,
    PriceResolution,
    ProductIndex,
)


def workbook_bytes() -> bytes:
    stream = BytesIO()
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet"
    worksheet.append(
        [
            "상품코드",
            "상품명",
            "규격",
            "현재고",
            "매입단가",
            "공급사코드",
            "공급사",
            "상품유형",
            "상품유형",
        ]
    )
    worksheet.append(["0001", "상품1", "1정", 3, 100, "S1", "공급사1", "A", "B"])
    worksheet.append(["0002", "상품2", "2정", -1, 99.5, "S2", "공급사2", "A", "B"])
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def build_test_client(tmp_path: Path):
    database_url = f"sqlite:///{tmp_path / 'jobs.db'}"
    settings = Settings(data_dir=tmp_path / "data", database_url=database_url)
    engine = build_engine(database_url)
    Base.metadata.create_all(engine)
    session_factory = build_session_factory(engine)
    app = create_app(settings)

    def override_get_db():
        with session_factory() as session:
            yield session

    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[get_settings] = lambda: settings
    return TestClient(app), session_factory, engine


def test_job_crud_and_excel_upload_share_job_read_contract(tmp_path: Path) -> None:
    client, session_factory, engine = build_test_client(tmp_path)
    content = workbook_bytes()

    with client:
        created = client.post("/api/jobs")
        job_id = created.json()["id"]
        uploaded = client.post(
            f"/api/jobs/{job_id}/excel",
            files={
                "file": (
                    "원본 상품리스트.xlsx",
                    content,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        detail = client.get(f"/api/jobs/{job_id}")
        listing = client.get("/api/jobs")

    assert created.status_code == 201
    assert uploaded.status_code == 200
    assert detail.status_code == 200
    assert listing.status_code == 200
    assert set(created.json()) == set(uploaded.json()) == set(detail.json())
    assert set(listing.json()[0]) == set(created.json())
    assert uploaded.json()["original_excel_name"] == "원본 상품리스트.xlsx"
    assert uploaded.json()["original_excel_sha256"] == hashlib.sha256(content).hexdigest()
    assert uploaded.json()["product_count"] == 2
    for response_body in (created.json(), uploaded.json(), detail.json(), listing.json()[0]):
        assert response_body["created_at"].endswith("Z")
        assert response_body["updated_at"].endswith("Z")

    with session_factory() as session:
        job = session.get(Job, job_id)
        products = session.scalars(
            select(ProductIndex)
            .where(ProductIndex.job_id == job_id)
            .order_by(ProductIndex.excel_row)
        ).all()
        stored_path = Path(job.original_excel_path)
        assert stored_path.name != "원본 상품리스트.xlsx"
        assert stored_path.suffix == ".xlsx"
        assert stored_path.exists()
        assert [product.product_code for product in products] == ["0001", "0002"]
        assert products[1].current_stock == -1
        assert products[1].purchase_price == 99.5

    engine.dispose()


def test_negative_purchase_price_upload_returns_validation_error(
    tmp_path: Path,
) -> None:
    client, _, engine = build_test_client(tmp_path)
    stream = BytesIO()
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet"
    worksheet.append(
        ["상품코드", "상품명", "규격", "현재고", "매입단가", "공급사코드", "공급사"]
    )
    worksheet.append(["0001", "상품", "1정", 0, -100, "S1", "공급사"])
    workbook.save(stream)
    workbook.close()

    with client:
        job_id = client.post("/api/jobs").json()["id"]
        response = client.post(
            f"/api/jobs/{job_id}/excel",
            files={"file": ("negative.xlsx", stream.getvalue())},
        )

    assert response.status_code == 422
    assert "매입단가는 0 이상" in response.json()["detail"]
    engine.dispose()


def test_failed_replacement_preserves_existing_file_and_index(tmp_path: Path) -> None:
    client, session_factory, engine = build_test_client(tmp_path)
    content = workbook_bytes()

    with client:
        job_id = client.post("/api/jobs").json()["id"]
        first = client.post(
            f"/api/jobs/{job_id}/excel",
            files={"file": ("products.xlsx", content)},
        )
        failed = client.post(
            f"/api/jobs/{job_id}/excel",
            files={"file": ("broken.xlsx", b"not an xlsx file")},
        )
        after = client.get(f"/api/jobs/{job_id}")

    assert first.status_code == 200
    assert failed.status_code == 422
    assert "손상되었거나 암호화된" in failed.json()["detail"]
    assert after.json()["original_excel_sha256"] == first.json()["original_excel_sha256"]
    assert after.json()["product_count"] == 2
    with session_factory() as session:
        job = session.get(Job, job_id)
        assert Path(job.original_excel_path).read_bytes() == content

    uploads = list((tmp_path / "data" / "uploads").iterdir())
    assert len(uploads) == 1
    engine.dispose()


def test_non_xlsx_upload_is_rejected_in_korean(tmp_path: Path) -> None:
    client, _, engine = build_test_client(tmp_path)
    with client:
        job_id = client.post("/api/jobs").json()["id"]
        response = client.post(
            f"/api/jobs/{job_id}/excel",
            files={"file": ("products.csv", b"code,name")},
        )

    assert response.status_code == 422
    assert response.json()["detail"] == (
        ".xlsx 형식의 상품리스트만 업로드할 수 있습니다."
    )
    engine.dispose()


def test_delete_job_removes_related_records_and_owned_files(tmp_path: Path) -> None:
    client, session_factory, engine = build_test_client(tmp_path)
    uploads_dir = tmp_path / "data" / "uploads"
    corrected_dir = tmp_path / "data" / "corrected"
    exports_dir = tmp_path / "data" / "exports"

    with client:
        job_id = client.post("/api/jobs").json()["id"]
        uploaded = client.post(
            f"/api/jobs/{job_id}/excel",
            files={"file": ("products.xlsx", workbook_bytes())},
        )
        assert uploaded.status_code == 200

        original_image_path = uploads_dir / "document-original.png"
        corrected_image_path = corrected_dir / "document-corrected.png"
        result_path = exports_dir / "result.xlsx"
        original_image_path.write_bytes(b"original")
        corrected_image_path.write_bytes(b"corrected")
        result_path.write_bytes(b"result")

        with session_factory() as session:
            job = session.get(Job, job_id)
            assert job is not None
            original_excel_path = Path(job.original_excel_path)
            document = Document(
                job_id=job_id,
                source_order=0,
                original_image_path=str(original_image_path),
                original_image_name="document.png",
                corrected_image_path=str(corrected_image_path),
                image_sha256="1" * 64,
            )
            session.add(document)
            session.flush()
            item = Item(
                document_id=document.id,
                source_row_order=0,
                product_name="상품",
            )
            session.add(item)
            session.flush()
            session.add(
                PriceResolution(
                    job_id=job_id,
                    product_code="0001",
                    selected_item_id=item.id,
                )
            )
            session.add(
                CompletedDocument(
                    job_id=job_id,
                    source_document_id=document.id,
                    image_sha256=document.image_sha256,
                )
            )
            job.status = JobStatus.COMPLETED
            job.result_path = str(result_path)
            session.commit()

        deleted = client.delete(f"/api/jobs/{job_id}")
        detail = client.get(f"/api/jobs/{job_id}")

    assert deleted.status_code == 204
    assert deleted.content == b""
    assert detail.status_code == 404
    for path in (
        original_excel_path,
        original_image_path,
        corrected_image_path,
        result_path,
    ):
        assert path.is_file() is False

    with session_factory() as session:
        assert session.get(Job, job_id) is None
        for model in (
            ProductIndex,
            Document,
            Item,
            PriceResolution,
            CompletedDocument,
        ):
            assert session.scalar(select(func.count(model.id))) == 0

    engine.dispose()


@pytest.mark.parametrize("job_status", [JobStatus.EXTRACTING, JobStatus.EXPORTING])
def test_delete_job_rejects_processing_status(
    tmp_path: Path, job_status: JobStatus
) -> None:
    client, session_factory, engine = build_test_client(tmp_path)

    with client:
        job_id = client.post("/api/jobs").json()["id"]
        with session_factory() as session:
            job = session.get(Job, job_id)
            assert job is not None
            job.status = job_status
            session.commit()

        response = client.delete(f"/api/jobs/{job_id}")
        detail = client.get(f"/api/jobs/{job_id}")

    assert response.status_code == 409
    assert response.json()["detail"] == (
        "추출 또는 내보내기 중인 작업은 삭제할 수 없습니다. 처리가 끝난 뒤 다시 시도해 주세요."
    )
    assert detail.status_code == 200
    assert detail.json()["status"] == job_status.value
    engine.dispose()
